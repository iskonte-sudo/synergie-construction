from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, UploadFile, File, Form, Query
from fastapi.responses import StreamingResponse, FileResponse
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from starlette.staticfiles import StaticFiles
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import io
import json
import uuid
import shutil
from pathlib import Path
from datetime import datetime, timedelta
from typing import List, Optional

from auth_utils import hash_password, verify_password, create_access_token, decode_token
from models import (
    UserCreate, UserUpdate, UserInDB, UserPublic, LoginRequest, LoginResponse, ChangePasswordRequest,
    QuoteCreate, Quote, QuoteStatusUpdate,
    MessageCreate, Message, MessageReply,
    ProjectCreate, Project,
    ServiceCreate, Service,
    SimulationCreate, Simulation,
    VisitCreate, Visit,
    Media, AuditLog,
    SiteSettings, SiteSettingsUpdate,
    SimulatorConfig, SimulatorConfigUpdate,
    SlideCreate, Slide,
    TestimonialCreate, Testimonial,
    FAQCreate, FAQ,
    TeamMemberCreate, TeamMember,
    PartnerCreate, Partner,
    BlogPostCreate, BlogPost,
    MenuItemCreate, MenuItem,
    ContentBlockUpdate, ContentBlock,
)

ROOT_DIR = Path(__file__).parent
UPLOADS_DIR = ROOT_DIR / 'uploads'
UPLOADS_DIR.mkdir(exist_ok=True)
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI(title='Synergie Construction API')
api_router = APIRouter(prefix='/api')
security = HTTPBearer(auto_error=False)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# Serve uploaded files
app.mount('/uploads', StaticFiles(directory=str(UPLOADS_DIR)), name='uploads')


# ----------------------- AUTH DEPENDENCIES -----------------------
async def get_current_user(credentials: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> UserInDB:
    if not credentials:
        raise HTTPException(status_code=401, detail='Non authentifié')
    payload = decode_token(credentials.credentials)
    if not payload or 'sub' not in payload:
        raise HTTPException(status_code=401, detail='Token invalide')
    user = await db.users.find_one({'id': payload['sub']})
    if not user:
        raise HTTPException(status_code=401, detail='Utilisateur introuvable')
    if not user.get('active', True):
        raise HTTPException(status_code=403, detail='Compte désactivé')
    user.pop('_id', None)
    return UserInDB(**user)


async def require_role(*roles: str):
    async def _check(user: UserInDB = Depends(get_current_user)) -> UserInDB:
        if user.role not in roles and user.role != 'super_admin':
            raise HTTPException(status_code=403, detail='Accès refusé')
        return user
    return _check


def require_super_admin(user: UserInDB = Depends(get_current_user)) -> UserInDB:
    if user.role != 'super_admin':
        raise HTTPException(status_code=403, detail='Super Admin requis')
    return user


def require_admin(user: UserInDB = Depends(get_current_user)) -> UserInDB:
    if user.role not in ('super_admin', 'admin'):
        raise HTTPException(status_code=403, detail='Administrateur requis')
    return user


async def log_action(user: UserInDB, action: str, target: str = None, details: dict = None, ip: str = None):
    log = AuditLog(user_id=user.id, user_email=user.email, action=action, target=target, details=details, ip=ip)
    await db.audit_logs.insert_one(log.model_dump())


def _clean(doc: dict) -> dict:
    if doc:
        doc.pop('_id', None)
    return doc


# ----------------------- AUTH ROUTES -----------------------
@api_router.post('/auth/login', response_model=LoginResponse)
async def login(payload: LoginRequest, request: Request):
    user = await db.users.find_one({'email': payload.email.lower()})
    if not user or not verify_password(payload.password, user.get('password_hash', '')):
        raise HTTPException(status_code=401, detail='Email ou mot de passe incorrect')
    if not user.get('active', True):
        raise HTTPException(status_code=403, detail='Compte désactivé')
    await db.users.update_one({'id': user['id']}, {'$set': {'last_login': datetime.utcnow()}})
    user['last_login'] = datetime.utcnow()
    token = create_access_token({'sub': user['id'], 'email': user['email'], 'role': user['role']})
    ip = request.client.host if request.client else None
    await db.audit_logs.insert_one(AuditLog(
        user_id=user['id'], user_email=user['email'], action='login', ip=ip
    ).model_dump())
    return LoginResponse(access_token=token, user=UserPublic(**user))


@api_router.get('/auth/me', response_model=UserPublic)
async def me(user: UserInDB = Depends(get_current_user)):
    return UserPublic(**user.model_dump())


@api_router.post('/auth/change-password')
async def change_password(payload: ChangePasswordRequest, user: UserInDB = Depends(get_current_user)):
    if not verify_password(payload.current_password, user.password_hash):
        raise HTTPException(status_code=400, detail='Mot de passe actuel incorrect')
    if len(payload.new_password) < 8:
        raise HTTPException(status_code=400, detail='Le nouveau mot de passe doit faire au moins 8 caractères')
    await db.users.update_one({'id': user.id}, {'$set': {'password_hash': hash_password(payload.new_password)}})
    await log_action(user, 'change_password')
    return {'ok': True}


# ----------------------- PUBLIC ROUTES -----------------------
@api_router.post('/public/quotes', response_model=Quote)
async def submit_quote(payload: QuoteCreate, request: Request):
    quote = Quote(**payload.model_dump())
    await db.quotes.insert_one(quote.model_dump())
    return quote


@api_router.post('/public/messages', response_model=Message)
async def submit_message(payload: MessageCreate):
    msg = Message(**payload.model_dump())
    await db.messages.insert_one(msg.model_dump())
    return msg


@api_router.post('/public/simulations', response_model=Simulation)
async def submit_simulation(payload: SimulationCreate):
    ref = 'SIM-' + str(int(datetime.utcnow().timestamp()))[-8:]
    sim = Simulation(**payload.model_dump(), reference=ref)
    await db.simulations.insert_one(sim.model_dump())
    return sim


@api_router.post('/public/visits')
async def track_visit(payload: VisitCreate, request: Request):
    ip = request.client.host if request.client else None
    visit = Visit(**payload.model_dump(), ip=ip)
    await db.visits.insert_one(visit.model_dump())
    return {'ok': True}


@api_router.get('/public/projects')
async def list_public_projects():
    items = await db.projects.find({'published': True}).sort('created_at', -1).to_list(500)
    return [_clean(i) for i in items]


@api_router.get('/public/services')
async def list_public_services():
    items = await db.services.find({'published': True}).sort('order', 1).to_list(100)
    return [_clean(i) for i in items]


@api_router.get('/public/settings')
async def public_settings():
    s = await db.settings.find_one({'id': 'main'})
    return _clean(s) if s else SiteSettings().model_dump()


@api_router.get('/')
async def root():
    return {'message': 'Synergie Construction API', 'version': '1.0'}


# ----------------------- ADMIN: DASHBOARD -----------------------
@api_router.get('/admin/dashboard')
async def admin_dashboard(user: UserInDB = Depends(get_current_user)):
    today = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)
    week_ago = today - timedelta(days=7)
    month_ago = today - timedelta(days=30)

    visits_total = await db.visits.count_documents({})
    visits_today = await db.visits.count_documents({'created_at': {'$gte': today}})
    quotes_total = await db.quotes.count_documents({})
    quotes_new = await db.quotes.count_documents({'status': 'nouveau'})
    messages_total = await db.messages.count_documents({})
    messages_new = await db.messages.count_documents({'status': 'nouveau'})
    projects_total = await db.projects.count_documents({})
    projects_published = await db.projects.count_documents({'published': True})
    simulations_total = await db.simulations.count_documents({})
    services_total = await db.services.count_documents({})

    # Monthly chart (last 6 months)
    months = []
    for i in range(5, -1, -1):
        m_start = (today.replace(day=1) - timedelta(days=30 * i)).replace(day=1, hour=0, minute=0, second=0)
        m_next = (m_start.replace(day=28) + timedelta(days=4)).replace(day=1)
        q_count = await db.quotes.count_documents({'created_at': {'$gte': m_start, '$lt': m_next}})
        v_count = await db.visits.count_documents({'created_at': {'$gte': m_start, '$lt': m_next}})
        s_count = await db.simulations.count_documents({'created_at': {'$gte': m_start, '$lt': m_next}})
        months.append({
            'month': m_start.strftime('%b %Y'),
            'quotes': q_count, 'visits': v_count, 'simulations': s_count,
        })

    # Top pages
    pipeline = [
        {'$group': {'_id': '$path', 'count': {'$sum': 1}}},
        {'$sort': {'count': -1}},
        {'$limit': 6},
    ]
    top_pages_raw = await db.visits.aggregate(pipeline).to_list(10)
    top_pages = [{'path': p['_id'] or '/', 'count': p['count']} for p in top_pages_raw]

    recent_quotes = await db.quotes.find().sort('created_at', -1).limit(5).to_list(5)
    recent_messages = await db.messages.find().sort('created_at', -1).limit(5).to_list(5)

    return {
        'stats': {
            'visits': {'total': visits_total, 'today': visits_today},
            'quotes': {'total': quotes_total, 'new': quotes_new},
            'messages': {'total': messages_total, 'new': messages_new},
            'projects': {'total': projects_total, 'published': projects_published},
            'simulations': {'total': simulations_total},
            'services': {'total': services_total},
        },
        'months': months,
        'top_pages': top_pages,
        'recent_quotes': [_clean(q) for q in recent_quotes],
        'recent_messages': [_clean(m) for m in recent_messages],
    }


# ----------------------- ADMIN: QUOTES -----------------------
@api_router.get('/admin/quotes')
async def list_quotes(
    user: UserInDB = Depends(get_current_user),
    status: Optional[str] = None,
    service: Optional[str] = None,
    q: Optional[str] = None,
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
):
    query = {}
    if status:
        query['status'] = status
    if service:
        query['service'] = service
    if q:
        query['$or'] = [
            {'name': {'$regex': q, '$options': 'i'}},
            {'email': {'$regex': q, '$options': 'i'}},
            {'phone': {'$regex': q, '$options': 'i'}},
        ]
    if from_date or to_date:
        rng = {}
        if from_date:
            rng['$gte'] = datetime.fromisoformat(from_date)
        if to_date:
            rng['$lte'] = datetime.fromisoformat(to_date)
        query['created_at'] = rng
    items = await db.quotes.find(query).sort('created_at', -1).to_list(2000)
    return [_clean(i) for i in items]


@api_router.get('/admin/quotes/export')
async def export_quotes(user: UserInDB = Depends(get_current_user)):
    from openpyxl import Workbook
    items = await db.quotes.find().sort('created_at', -1).to_list(5000)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Devis'
    ws.append(['Date', 'Référence', 'Service', 'Nom', 'Email', 'Téléphone', 'Statut', 'Détails'])
    for q in items:
        ws.append([
            q.get('created_at').strftime('%Y-%m-%d %H:%M') if q.get('created_at') else '',
            q.get('id', '')[:8],
            q.get('service_title') or q.get('service', ''),
            q.get('name', ''),
            q.get('email', ''),
            q.get('phone', ''),
            q.get('status', ''),
            json.dumps(q.get('values', {}), ensure_ascii=False),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await log_action(user, 'export_quotes')
    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=devis-' + datetime.utcnow().strftime('%Y%m%d') + '.xlsx'}
    )


@api_router.patch('/admin/quotes/{quote_id}')
async def update_quote(quote_id: str, payload: QuoteStatusUpdate, user: UserInDB = Depends(get_current_user)):
    upd = {k: v for k, v in payload.model_dump().items() if v is not None}
    upd['updated_at'] = datetime.utcnow()
    res = await db.quotes.update_one({'id': quote_id}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Devis introuvable')
    await log_action(user, 'update_quote', quote_id, upd)
    doc = await db.quotes.find_one({'id': quote_id})
    return _clean(doc)


@api_router.delete('/admin/quotes/{quote_id}')
async def delete_quote(quote_id: str, user: UserInDB = Depends(require_admin)):
    await db.quotes.delete_one({'id': quote_id})
    await log_action(user, 'delete_quote', quote_id)
    return {'ok': True}


# ----------------------- ADMIN: MESSAGES -----------------------
@api_router.get('/admin/messages')
async def list_messages(user: UserInDB = Depends(get_current_user), status: Optional[str] = None, q: Optional[str] = None):
    query = {}
    if status:
        query['status'] = status
    if q:
        query['$or'] = [
            {'name': {'$regex': q, '$options': 'i'}},
            {'email': {'$regex': q, '$options': 'i'}},
            {'message': {'$regex': q, '$options': 'i'}},
        ]
    items = await db.messages.find(query).sort('created_at', -1).to_list(2000)
    return [_clean(i) for i in items]


@api_router.patch('/admin/messages/{msg_id}')
async def update_message(msg_id: str, payload: MessageReply, user: UserInDB = Depends(get_current_user)):
    upd = {'status': payload.status}
    if payload.reply:
        upd['reply'] = payload.reply
        upd['replied_at'] = datetime.utcnow()
    res = await db.messages.update_one({'id': msg_id}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Message introuvable')
    await log_action(user, 'reply_message', msg_id)
    return _clean(await db.messages.find_one({'id': msg_id}))


@api_router.delete('/admin/messages/{msg_id}')
async def delete_message(msg_id: str, user: UserInDB = Depends(require_admin)):
    await db.messages.delete_one({'id': msg_id})
    await log_action(user, 'delete_message', msg_id)
    return {'ok': True}


# ----------------------- ADMIN: PROJECTS -----------------------
@api_router.get('/admin/projects')
async def admin_list_projects(user: UserInDB = Depends(get_current_user)):
    items = await db.projects.find().sort('created_at', -1).to_list(1000)
    return [_clean(i) for i in items]


@api_router.post('/admin/projects', response_model=Project)
async def create_project(payload: ProjectCreate, user: UserInDB = Depends(get_current_user)):
    p = Project(**payload.model_dump())
    await db.projects.insert_one(p.model_dump())
    await log_action(user, 'create_project', p.id)
    return p


@api_router.patch('/admin/projects/{pid}')
async def update_project(pid: str, payload: ProjectCreate, user: UserInDB = Depends(get_current_user)):
    upd = payload.model_dump()
    upd['updated_at'] = datetime.utcnow()
    res = await db.projects.update_one({'id': pid}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Projet introuvable')
    await log_action(user, 'update_project', pid)
    return _clean(await db.projects.find_one({'id': pid}))


@api_router.delete('/admin/projects/{pid}')
async def delete_project(pid: str, user: UserInDB = Depends(require_admin)):
    await db.projects.delete_one({'id': pid})
    await log_action(user, 'delete_project', pid)
    return {'ok': True}


# ----------------------- ADMIN: SERVICES -----------------------
@api_router.get('/admin/services')
async def admin_list_services(user: UserInDB = Depends(get_current_user)):
    items = await db.services.find().sort('order', 1).to_list(200)
    return [_clean(i) for i in items]


@api_router.post('/admin/services', response_model=Service)
async def create_service(payload: ServiceCreate, user: UserInDB = Depends(get_current_user)):
    s = Service(**payload.model_dump())
    await db.services.insert_one(s.model_dump())
    await log_action(user, 'create_service', s.id)
    return s


@api_router.patch('/admin/services/{sid}')
async def update_service(sid: str, payload: ServiceCreate, user: UserInDB = Depends(get_current_user)):
    upd = payload.model_dump()
    upd['updated_at'] = datetime.utcnow()
    res = await db.services.update_one({'id': sid}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Service introuvable')
    await log_action(user, 'update_service', sid)
    return _clean(await db.services.find_one({'id': sid}))


@api_router.delete('/admin/services/{sid}')
async def delete_service(sid: str, user: UserInDB = Depends(require_admin)):
    await db.services.delete_one({'id': sid})
    await log_action(user, 'delete_service', sid)
    return {'ok': True}


# ----------------------- ADMIN: SIMULATIONS -----------------------
@api_router.get('/admin/simulations')
async def list_simulations(user: UserInDB = Depends(get_current_user)):
    items = await db.simulations.find().sort('created_at', -1).to_list(2000)
    return [_clean(i) for i in items]


@api_router.get('/admin/simulations/export')
async def export_simulations(user: UserInDB = Depends(get_current_user)):
    from openpyxl import Workbook
    items = await db.simulations.find().sort('created_at', -1).to_list(5000)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Simulations'
    ws.append(['Date', 'Référence', 'Type', 'Surface', 'Budget annoncé', 'Délai', 'Estimation basse', 'Estimation haute', 'Nom', 'Email', 'Téléphone'])
    for s in items:
        c = s.get('contact', {})
        ws.append([
            s.get('created_at').strftime('%Y-%m-%d %H:%M') if s.get('created_at') else '',
            s.get('reference', ''),
            s.get('project_type', ''),
            s.get('surface', ''),
            s.get('budget', ''),
            s.get('delai', ''),
            s.get('estimate_low', 0),
            s.get('estimate_high', 0),
            c.get('name', ''),
            c.get('email', ''),
            c.get('phone', ''),
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await log_action(user, 'export_simulations')
    return StreamingResponse(
        buf,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename=simulations-' + datetime.utcnow().strftime('%Y%m%d') + '.xlsx'}
    )


@api_router.delete('/admin/simulations/{sid}')
async def delete_simulation(sid: str, user: UserInDB = Depends(require_admin)):
    await db.simulations.delete_one({'id': sid})
    await log_action(user, 'delete_simulation', sid)
    return {'ok': True}


# ----------------------- ADMIN: USERS -----------------------
@api_router.get('/admin/users')
async def list_users(user: UserInDB = Depends(require_super_admin)):
    items = await db.users.find().sort('created_at', -1).to_list(200)
    return [{k: v for k, v in _clean(u).items() if k != 'password_hash'} for u in items]


@api_router.post('/admin/users')
async def create_user(payload: UserCreate, user: UserInDB = Depends(require_super_admin)):
    payload.email = payload.email.lower()
    existing = await db.users.find_one({'email': payload.email})
    if existing:
        raise HTTPException(400, 'Email déjà utilisé')
    if len(payload.password) < 8:
        raise HTTPException(400, 'Mot de passe trop court (8 minimum)')
    if payload.role not in ('super_admin', 'admin', 'editor'):
        raise HTTPException(400, 'Rôle invalide')
    u = UserInDB(**payload.model_dump(exclude={'password'}), password_hash=hash_password(payload.password))
    await db.users.insert_one(u.model_dump())
    await log_action(user, 'create_user', u.id, {'email': u.email})
    return UserPublic(**u.model_dump())


@api_router.patch('/admin/users/{uid}')
async def update_user(uid: str, payload: UserUpdate, user: UserInDB = Depends(require_super_admin)):
    upd = {k: v for k, v in payload.model_dump().items() if v is not None}
    if 'password' in upd:
        if len(upd['password']) < 8:
            raise HTTPException(400, 'Mot de passe trop court')
        upd['password_hash'] = hash_password(upd.pop('password'))
    if 'role' in upd and upd['role'] not in ('super_admin', 'admin', 'editor'):
        raise HTTPException(400, 'Rôle invalide')
    res = await db.users.update_one({'id': uid}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Utilisateur introuvable')
    await log_action(user, 'update_user', uid)
    doc = await db.users.find_one({'id': uid})
    doc = _clean(doc)
    doc.pop('password_hash', None)
    return doc


@api_router.delete('/admin/users/{uid}')
async def delete_user(uid: str, user: UserInDB = Depends(require_super_admin)):
    if uid == user.id:
        raise HTTPException(400, 'Vous ne pouvez pas vous supprimer vous-même')
    await db.users.delete_one({'id': uid})
    await log_action(user, 'delete_user', uid)
    return {'ok': True}


# ----------------------- ADMIN: SETTINGS -----------------------
@api_router.get('/admin/settings')
async def get_settings(user: UserInDB = Depends(get_current_user)):
    s = await db.settings.find_one({'id': 'main'})
    if not s:
        s = SiteSettings().model_dump()
        await db.settings.insert_one(s)
    return _clean(s)


@api_router.put('/admin/settings')
async def update_settings(payload: SiteSettingsUpdate, user: UserInDB = Depends(require_admin)):
    upd = {k: v for k, v in payload.model_dump().items() if v is not None}
    upd['updated_at'] = datetime.utcnow()
    await db.settings.update_one({'id': 'main'}, {'$set': upd}, upsert=True)
    await log_action(user, 'update_settings')
    s = await db.settings.find_one({'id': 'main'})
    return _clean(s)


# ----------------------- ADMIN: MEDIA -----------------------
@api_router.post('/admin/media')
async def upload_media(
    file: UploadFile = File(...),
    folder: str = Form('general'),
    user: UserInDB = Depends(get_current_user),
):
    ext = Path(file.filename).suffix
    name = f"{uuid.uuid4().hex}{ext}"
    folder_path = UPLOADS_DIR / folder
    folder_path.mkdir(exist_ok=True, parents=True)
    file_path = folder_path / name
    with open(file_path, 'wb') as f:
        shutil.copyfileobj(file.file, f)
    size = file_path.stat().st_size
    url = f"/uploads/{folder}/{name}"
    media = Media(
        name=file.filename, url=url, mime=file.content_type or 'application/octet-stream',
        size=size, folder=folder, uploaded_by=user.email,
    )
    await db.media.insert_one(media.model_dump())
    return media


@api_router.get('/admin/media')
async def list_media(user: UserInDB = Depends(get_current_user), folder: Optional[str] = None, q: Optional[str] = None):
    query = {}
    if folder:
        query['folder'] = folder
    if q:
        query['name'] = {'$regex': q, '$options': 'i'}
    items = await db.media.find(query).sort('created_at', -1).to_list(500)
    return [_clean(i) for i in items]


@api_router.delete('/admin/media/{mid}')
async def delete_media(mid: str, user: UserInDB = Depends(get_current_user)):
    m = await db.media.find_one({'id': mid})
    if not m:
        raise HTTPException(404, 'Média introuvable')
    try:
        p = UPLOADS_DIR / m['folder'] / Path(m['url']).name
        if p.exists():
            p.unlink()
    except Exception as e:
        logger.warning(f"Failed to delete file: {e}")
    await db.media.delete_one({'id': mid})
    await log_action(user, 'delete_media', mid)
    return {'ok': True}


# ----------------------- ADMIN: AUDIT LOGS -----------------------
@api_router.get('/admin/audit-logs')
async def list_audit_logs(user: UserInDB = Depends(require_admin), limit: int = 200):
    items = await db.audit_logs.find().sort('created_at', -1).limit(limit).to_list(limit)
    return [_clean(i) for i in items]


# ----------------------- SIMULATOR CONFIG -----------------------
@api_router.get('/public/simulator-config')
async def public_simulator_config():
    s = await db.simulator_config.find_one({'id': 'main'})
    return _clean(s) if s else SimulatorConfig().model_dump()


@api_router.get('/admin/simulator-config')
async def get_simulator_config(user: UserInDB = Depends(get_current_user)):
    s = await db.simulator_config.find_one({'id': 'main'})
    if not s:
        s = SimulatorConfig().model_dump()
        await db.simulator_config.insert_one(s)
    return _clean(s)


@api_router.put('/admin/simulator-config')
async def update_simulator_config(payload: SimulatorConfigUpdate, user: UserInDB = Depends(require_admin)):
    upd = {k: v for k, v in payload.model_dump().items() if v is not None}
    upd['updated_at'] = datetime.utcnow()
    await db.simulator_config.update_one({'id': 'main'}, {'$set': upd}, upsert=True)
    await log_action(user, 'update_simulator_config')
    s = await db.simulator_config.find_one({'id': 'main'})
    return _clean(s)


# ----------------------- GENERIC CRUD HELPERS -----------------------
def make_crud(collection: str, ModelIn, ModelOut, prefix: str, public_filter: dict = None):
    """Registers CRUD endpoints for a collection under admin & public paths."""
    @api_router.get(f'/public/{prefix}')
    async def _public_list():
        query = public_filter or {'active': True}
        items = await db[collection].find(query).sort('order', 1).to_list(500)
        return [_clean(i) for i in items]

    _public_list.__name__ = f'public_list_{prefix}'

    @api_router.get(f'/admin/{prefix}')
    async def _admin_list(user: UserInDB = Depends(get_current_user)):
        items = await db[collection].find().sort('order', 1).to_list(1000)
        return [_clean(i) for i in items]

    _admin_list.__name__ = f'admin_list_{prefix}'

    @api_router.post(f'/admin/{prefix}')
    async def _create(payload: ModelIn, user: UserInDB = Depends(get_current_user)):
        obj = ModelOut(**payload.model_dump())
        await db[collection].insert_one(obj.model_dump())
        await log_action(user, f'create_{prefix}', obj.id)
        return obj

    _create.__name__ = f'create_{prefix}'

    @api_router.patch(f'/admin/{prefix}/{{item_id}}')
    async def _update(item_id: str, payload: ModelIn, user: UserInDB = Depends(get_current_user)):
        upd = payload.model_dump()
        upd['updated_at'] = datetime.utcnow()
        res = await db[collection].update_one({'id': item_id}, {'$set': upd})
        if res.matched_count == 0:
            raise HTTPException(404, 'Introuvable')
        await log_action(user, f'update_{prefix}', item_id)
        return _clean(await db[collection].find_one({'id': item_id}))

    _update.__name__ = f'update_{prefix}'

    @api_router.delete(f'/admin/{prefix}/{{item_id}}')
    async def _delete(item_id: str, user: UserInDB = Depends(require_admin)):
        await db[collection].delete_one({'id': item_id})
        await log_action(user, f'delete_{prefix}', item_id)
        return {'ok': True}

    _delete.__name__ = f'delete_{prefix}'


# Register CRUDs
make_crud('slides', SlideCreate, Slide, 'slides')
make_crud('testimonials', TestimonialCreate, Testimonial, 'testimonials')
make_crud('faqs', FAQCreate, FAQ, 'faqs')
make_crud('team', TeamMemberCreate, TeamMember, 'team')
make_crud('partners', PartnerCreate, Partner, 'partners')
make_crud('menu_items', MenuItemCreate, MenuItem, 'menu-items')


# ----------------------- BLOG -----------------------
@api_router.get('/public/blog')
async def public_blog_list():
    items = await db.blog_posts.find({'published': True}).sort('published_at', -1).to_list(500)
    return [_clean(i) for i in items]


@api_router.get('/public/blog/{slug}')
async def public_blog_detail(slug: str):
    item = await db.blog_posts.find_one({'slug': slug, 'published': True})
    if not item:
        raise HTTPException(404, 'Article introuvable')
    await db.blog_posts.update_one({'id': item['id']}, {'$inc': {'views': 1}})
    return _clean(item)


@api_router.get('/admin/blog')
async def admin_blog_list(user: UserInDB = Depends(get_current_user)):
    items = await db.blog_posts.find().sort('created_at', -1).to_list(1000)
    return [_clean(i) for i in items]


@api_router.post('/admin/blog')
async def create_blog(payload: BlogPostCreate, user: UserInDB = Depends(get_current_user)):
    if await db.blog_posts.find_one({'slug': payload.slug}):
        raise HTTPException(400, 'Ce slug est déjà utilisé')
    post = BlogPost(**payload.model_dump())
    if post.published and not post.published_at:
        post.published_at = datetime.utcnow()
    await db.blog_posts.insert_one(post.model_dump())
    await log_action(user, 'create_blog', post.id)
    return post


@api_router.patch('/admin/blog/{item_id}')
async def update_blog(item_id: str, payload: BlogPostCreate, user: UserInDB = Depends(get_current_user)):
    upd = payload.model_dump()
    upd['updated_at'] = datetime.utcnow()
    if upd.get('published') and not upd.get('published_at'):
        upd['published_at'] = datetime.utcnow()
    res = await db.blog_posts.update_one({'id': item_id}, {'$set': upd})
    if res.matched_count == 0:
        raise HTTPException(404, 'Article introuvable')
    await log_action(user, 'update_blog', item_id)
    return _clean(await db.blog_posts.find_one({'id': item_id}))


@api_router.delete('/admin/blog/{item_id}')
async def delete_blog(item_id: str, user: UserInDB = Depends(require_admin)):
    await db.blog_posts.delete_one({'id': item_id})
    await log_action(user, 'delete_blog', item_id)
    return {'ok': True}


# ----------------------- CONTENT BLOCKS (key/value editable text) -----------------------
@api_router.get('/public/content')
async def public_content(page: Optional[str] = None):
    q = {}
    if page:
        q['page'] = page
    items = await db.content_blocks.find(q).to_list(2000)
    return {i['key']: i['value'] for i in items}


@api_router.get('/admin/content')
async def admin_content(user: UserInDB = Depends(get_current_user), page: Optional[str] = None):
    q = {}
    if page:
        q['page'] = page
    items = await db.content_blocks.find(q).to_list(2000)
    return [_clean(i) for i in items]


@api_router.put('/admin/content')
async def upsert_content(payload: ContentBlockUpdate, user: UserInDB = Depends(get_current_user)):
    existing = await db.content_blocks.find_one({'key': payload.key})
    doc = {
        'key': payload.key,
        'value': payload.value,
        'page': payload.page or 'general',
        'label': payload.label or '',
        'updated_at': datetime.utcnow(),
    }
    if existing:
        await db.content_blocks.update_one({'key': payload.key}, {'$set': doc})
    else:
        doc['id'] = str(uuid.uuid4())
        await db.content_blocks.insert_one(doc)
    await log_action(user, 'update_content', payload.key)
    return {'ok': True, 'key': payload.key}


@api_router.delete('/admin/content/{key}')
async def delete_content(key: str, user: UserInDB = Depends(require_admin)):
    await db.content_blocks.delete_one({'key': key})
    await log_action(user, 'delete_content', key)
    return {'ok': True}


# ----------------------- APP SETUP -----------------------
app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=['*'],
    allow_methods=['*'],
    allow_headers=['*'],
)


@app.on_event('startup')
async def startup_event():
    # Seed super admin
    existing = await db.users.find_one({'email': 'admin@synergieconstruction.com'})
    if not existing:
        admin = UserInDB(
            email='admin@synergieconstruction.com',
            name='Super Administrateur',
            role='super_admin',
            password_hash=hash_password('Synergie2025!'),
        )
        await db.users.insert_one(admin.model_dump())
        logger.info('Super admin seeded: admin@synergieconstruction.com')

    # Seed settings
    existing_settings = await db.settings.find_one({'id': 'main'})
    if not existing_settings:
        s = SiteSettings(socials=[
            {'name': 'Facebook', 'icon': 'Facebook', 'url': 'https://www.facebook.com/synergieconstruction'},
            {'name': 'Instagram', 'icon': 'Instagram', 'url': 'https://www.instagram.com/synergieconstruction'},
            {'name': 'LinkedIn', 'icon': 'Linkedin', 'url': 'https://www.linkedin.com/company/synergieconstruction'},
            {'name': 'WhatsApp', 'icon': 'MessageCircle', 'url': 'https://wa.me/221771658042'},
            {'name': 'YouTube', 'icon': 'Youtube', 'url': 'https://www.youtube.com/@synergieconstruction'},
            {'name': 'TikTok', 'icon': 'Music2', 'url': 'https://www.tiktok.com/@synergieconstruction'},
        ])
        await db.settings.insert_one(s.model_dump())
        logger.info('Site settings seeded')


    # Seed simulator config
    existing_sim = await db.simulator_config.find_one({'id': 'main'})
    if not existing_sim:
        sim_cfg = SimulatorConfig(
            project_types=[
                {'id': 'villa', 'label': 'Villa', 'icon': 'Home', 'baseCostPerSqm': 280000, 'monthsPerSqm': 0.012},
                {'id': 'immeuble', 'label': 'Immeuble', 'icon': 'Building2', 'baseCostPerSqm': 350000, 'monthsPerSqm': 0.015},
                {'id': 'bureau', 'label': 'Bureau', 'icon': 'Briefcase', 'baseCostPerSqm': 320000, 'monthsPerSqm': 0.013},
                {'id': 'commerce', 'label': 'Commerce', 'icon': 'Store', 'baseCostPerSqm': 300000, 'monthsPerSqm': 0.011},
                {'id': 'hotel', 'label': 'Hôtel', 'icon': 'Hotel', 'baseCostPerSqm': 420000, 'monthsPerSqm': 0.018},
                {'id': 'entrepot', 'label': 'Entrepôt', 'icon': 'Warehouse', 'baseCostPerSqm': 180000, 'monthsPerSqm': 0.008},
                {'id': 'renovation', 'label': 'Rénovation', 'icon': 'Wrench', 'baseCostPerSqm': 150000, 'monthsPerSqm': 0.009},
            ],
            surface_options=[
                {'id': 's1', 'label': 'Moins de 100 m²', 'value': 80},
                {'id': 's2', 'label': '100 à 250 m²', 'value': 175},
                {'id': 's3', 'label': '250 à 500 m²', 'value': 375},
                {'id': 's4', 'label': '500 à 1000 m²', 'value': 750},
                {'id': 's5', 'label': 'Plus de 1000 m²', 'value': 1500},
            ],
            prestation_options=[
                {'id': 'faisabilite', 'label': 'Étude de faisabilité', 'icon': 'Search', 'recommends': 'conseil-technique'},
                {'id': 'plans-archi', 'label': 'Plans architecturaux', 'icon': 'PenTool', 'recommends': 'plans-2d-3d'},
                {'id': 'modelisation-3d', 'label': 'Modélisation 3D', 'icon': 'Box', 'recommends': 'plans-2d-3d'},
                {'id': 'rendus', 'label': 'Rendus photoréalistes', 'icon': 'Image', 'recommends': 'plans-2d-3d'},
                {'id': 'controle-qualite', 'label': 'Contrôle qualité', 'icon': 'ShieldCheck', 'recommends': 'suivi-controle'},
                {'id': 'suivi-chantier', 'label': 'Suivi de chantier', 'icon': 'HardHat', 'recommends': 'suivi-controle'},
                {'id': 'coordination', 'label': 'Coordination des intervenants', 'icon': 'Users', 'recommends': 'suivi-controle'},
                {'id': 'amoa', 'label': "Assistance maîtrise d'ouvrage", 'icon': 'ClipboardCheck', 'recommends': 'conseil-technique'},
            ],
            budget_options=[
                {'id': 'b1', 'label': 'Moins de 25 millions FCFA', 'min': 0, 'max': 25},
                {'id': 'b2', 'label': '25 à 50 millions FCFA', 'min': 25, 'max': 50},
                {'id': 'b3', 'label': '50 à 100 millions FCFA', 'min': 50, 'max': 100},
                {'id': 'b4', 'label': 'Plus de 100 millions FCFA', 'min': 100, 'max': 999},
            ],
            delai_options=[
                {'id': 'd1', 'label': 'Urgent (moins de 3 mois)', 'months': 2},
                {'id': 'd2', 'label': '3 à 6 mois', 'months': 5},
                {'id': 'd3', 'label': '6 à 12 mois', 'months': 9},
                {'id': 'd4', 'label': '12 à 24 mois', 'months': 18},
                {'id': 'd5', 'label': 'Plus de 24 mois', 'months': 28},
                {'id': 'd6', 'label': 'Flexible', 'months': 12},
            ],
        )
        await db.simulator_config.insert_one(sim_cfg.model_dump())
        logger.info('Simulator config seeded')

    # Seed slides
    if await db.slides.count_documents({}) == 0:
        for i, s in enumerate([
            {'title': "Transformer vos projets en réalité", 'subtitle': "Excellence et innovation dans chaque construction", 'image': 'https://images.unsplash.com/photo-1541888946425-d81bb19240f5?auto=format&fit=crop&w=1920&q=80', 'button_label': 'Nos services', 'button_link': '/services'},
            {'title': "Construire l'avenir avec précision", 'subtitle': "Des fondations solides pour des projets durables", 'image': 'https://images.unsplash.com/photo-1508450859948-4e04fabaa4ea?auto=format&fit=crop&w=1920&q=80', 'button_label': 'Nos réalisations', 'button_link': '/realisations'},
            {'title': "L'expertise au service de vos projets", 'subtitle': "30 ans d'expérience dans la construction", 'image': 'https://images.unsplash.com/photo-1580587771525-78b9dba3b914?auto=format&fit=crop&w=1920&q=80', 'button_label': 'Contactez-nous', 'button_link': '/contact'},
        ]):
            await db.slides.insert_one(Slide(**s, order=i).model_dump())
        logger.info('Slides seeded')

    # Seed testimonials
    if await db.testimonials.count_documents({}) == 0:
        for i, t in enumerate([
            {'name': 'Mamadou Diop', 'role': 'Propriétaire', 'content': "Synergie Construction a transformé notre projet en réalité. Leur professionnalisme et leur rigueur sont remarquables. Je recommande vivement.", 'rating': 5, 'image': 'https://i.pravatar.cc/150?img=12'},
            {'name': 'Aïssatou Ndiaye', 'role': 'Directrice, Ndiaye & Co', 'content': "Notre nouvel immeuble de bureaux est exactement ce que nous voulions. Travail livré dans les délais et budget respecté.", 'rating': 5, 'image': 'https://i.pravatar.cc/150?img=44'},
            {'name': 'Cheikh Fall', 'role': 'Promoteur Immobilier', 'content': "Une équipe d'experts, à l'écoute et toujours disponibles. Plusieurs projets réalisés ensemble, toujours avec excellence.", 'rating': 5, 'image': 'https://i.pravatar.cc/150?img=33'},
            {'name': 'Fatou Sarr', 'role': 'Architecte', 'content': "Collaboration exceptionnelle. Les plans 3D ont permis à mes clients de visualiser parfaitement leur projet.", 'rating': 5, 'image': 'https://i.pravatar.cc/150?img=47'},
        ]):
            await db.testimonials.insert_one(Testimonial(**t, order=i).model_dump())
        logger.info('Testimonials seeded')

    # Seed FAQs
    if await db.faqs.count_documents({}) == 0:
        for i, f in enumerate([
            {'question': 'Quels types de projets prenez-vous en charge ?', 'answer': "Nous prenons en charge tous types de projets : résidentiels (villas, immeubles), commerciaux, industriels et institutionnels. De la conception à la livraison clé en main."},
            {'question': 'Comment se déroule la demande de devis ?', 'answer': "Remplissez notre formulaire en ligne ou contactez-nous par WhatsApp. Un expert vous rappelle sous 24h pour étudier votre projet et établir un devis détaillé gratuit."},
            {'question': "Quels sont vos délais d'intervention ?", 'answer': "Les délais varient selon la nature et l'ampleur du projet. Une étude préliminaire est réalisée sous 48h, puis un planning détaillé est établi avec le client."},
            {'question': 'Proposez-vous des financements ?', 'answer': "Oui, nous travaillons avec des partenaires financiers pour faciliter le financement de vos projets. Échelonnement possible selon les étapes du chantier."},
            {'question': 'Quelles garanties offrez-vous ?', 'answer': "Tous nos travaux sont couverts par une garantie décennale. Nous garantissons également la qualité des matériaux et la conformité aux normes en vigueur."},
            {'question': 'Intervenez-vous en dehors de Dakar ?', 'answer': "Oui, nous intervenons dans tout le Sénégal (Thiès, Saint-Louis, Saly, Mbour, Touba, Ziguinchor...) et dans la sous-région ouest-africaine selon les projets."},
        ]):
            await db.faqs.insert_one(FAQ(**f, order=i).model_dump())
        logger.info('FAQs seeded')

    # Seed menu items (header)
    if await db.menu_items.count_documents({}) == 0:
        for i, m in enumerate([
            {'label': 'Accueil', 'path': '/', 'location': 'header'},
            {'label': 'À propos', 'path': '/a-propos', 'location': 'header'},
            {'label': 'Services', 'path': '/services', 'location': 'header'},
            {'label': 'Nos Réalisations', 'path': '/realisations', 'location': 'header'},
            {'label': 'Blog', 'path': '/blog', 'location': 'header'},
            {'label': 'Contact', 'path': '/contact', 'location': 'header'},
        ]):
            await db.menu_items.insert_one(MenuItem(**m, order=i).model_dump())
        logger.info('Menu items seeded')


@app.on_event('shutdown')
async def shutdown_db_client():
    client.close()
